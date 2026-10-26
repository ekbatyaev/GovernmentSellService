import asyncio
import os
import pandas as pd
from datetime import datetime, timedelta, date
from fastapi import HTTPException, BackgroundTasks
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Side, Border
from openpyxl.utils import get_column_letter
from app.backend.database.static_info import REGIONS_OF_THE_FILTERS, ALL_REGION_CODES
from app.backend.parsers.request_archives import get_docs_by_region, download_archive_from_result
from app.backend.parsers.xml_parser import parse_zip_archive_purchases, parse_zip_archive_protocols
from app.settings import settings, logger
from typing import Dict
from app.backend.email.functions import send_email
from app.backend.api_client import api_datum_query
from app.backend.parsers.schemas import Notice

# Сколько регионов обрабатывать одновременно.
REGION_CONCURRENCY = 4
# Сколько архивов (суммарно по всем регионам) скачивать и парсить одновременно.
ARCHIVE_CONCURRENCY = 8
# Сколько записей (purchase/protocol, суммарно по всем архивам) одновременно
# писать через api_datum_query.
RECORD_CONCURRENCY = 10

# Состояние текущей фоновой задачи админки. status: idle | running | done | error
current_job: dict = {"name": None, "status": "idle", "message": None}

async def run_job(name: str, func, *args, **kwargs):
    global current_job
    try:
        await func(*args, **kwargs)
        current_job = {"name": name, "status": "done", "message": f"{name}: завершено"}
    except Exception as e:
        logger.exception("Фоновая задача %s упала", name)
        current_job = {"name": name, "status": "error", "message": f"{name}: {e}"}


def start_job(background: BackgroundTasks, name: str, func, *args, **kwargs):
    global current_job
    if current_job["status"] == "running":
        raise HTTPException(status_code=409, detail=f"Уже выполняется задача {current_job['name']}, пожалуйста, попробуйте позже.")
    current_job = {"name": name, "status": "running", "message": None}
    background.add_task(run_job, name, func, *args, **kwargs)

def _build_analysis_workbook(rows, analysis_path) -> None:
    """Синхронная сборка Excel-файла (была телом create_analysis).
    Вызывается через asyncio.to_thread, чтобы не блокировать event loop."""

    rows.sort(
        key=lambda r: r["Дата публикации"] or datetime.min,
        reverse=True,
    )

    rows = rows[:50000]

    df = pd.DataFrame(rows)
    df.to_excel(analysis_path, index=False)

    wb = load_workbook(analysis_path)
    ws = wb.active

    header_fill = PatternFill("solid", fgColor="366092")
    header_font = Font(color="FFFFFF", bold=True)
    cell_font = Font(size=10)
    align = Alignment(wrap_text=True, vertical="center")
    thin_side = Side(style="thin")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for col in range(1, ws.max_column + 1):
        c = ws.cell(1, col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align
        c.border = border

    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 30
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = cell_font
            cell.alignment = align
            cell.border = border

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = max(
            len(str(ws.cell(r, col).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[letter].width = min(max_len + 2, 50)

    last = ws.max_row + 2
    ws.cell(last, 1, "Данные по закупкам, лотам и позициям")
    ws.merge_cells(
        start_row=last,
        start_column=1,
        end_row=last,
        end_column=ws.max_column,
    )

    footnote_cell = ws.cell(last, 1)
    footnote_cell.font = Font(italic=True, size=9, color="555555")
    footnote_cell.alignment = Alignment(horizontal="center")
    footnote_cell.border = Border(top=Side(style="thin", color="AAAAAA"))

    wb.save(analysis_path)


async def create_analysis(rows, analysis_path):
    try:
        await asyncio.to_thread(_build_analysis_workbook, rows, analysis_path)
    except Exception:
        logger.exception("Pipeline: ошибка при создании анализа | path=%s", analysis_path)


async def send_analysis(date_of_processing_notices: date, rows: list, emails: list, created: int, updated: int, skipped: int, extra_rows: list = None, path_name_extra="all_purchases.xlsx") -> None:
    date_str = date_of_processing_notices.strftime('%d.%m.%Y')

    logger.info(f"""При обработке закупок за {date_str}\n
    Новых заявок добавлено: {created}\n
    Заявок обновлено: {updated}\n
    Пропущено: {skipped}
    """)

    if created == 0 and updated == 0 and skipped == 0:
        return

    html_content = f"""
                    <html><body style="font-family:Arial;">
                    <h2 style="color:#2E86C1;">Уведомление о заявках с госзакупок</h2>
                    <p>Новых заявок добавлено: <b style="color:#E74C3C;font-size:18px;">{created}</b></p>
                    <p>Заявок обновлено: <b style="color:#F39C12;font-size:18px;">{updated}</b></p>
                    <p>Пропущено: <b style="color:#7F8C8D;font-size:18px;">{skipped}</b></p>
                    <hr><p style="color:#888;font-size:12px;">Это письмо сформировано автоматически, отвечать на него не нужно. 
                    В excel-таблицах отображается максимум 50 000 строк-заявок.</p>
                    </body></html>
                    """

    attachments = []

    analysis_path = "analysis.xlsx"

    await create_analysis(rows, analysis_path=analysis_path)

    attachments.append(analysis_path)

    if extra_rows:

        await create_analysis(extra_rows, analysis_path=path_name_extra)
        attachments.append(path_name_extra)

    subject = f"Заявки с госзакупок за {date_str}"

    try:

        for email in emails:
            try:
                await send_email(
                    email,
                    subject,
                    html_content,
                    attachments=attachments if (created or updated) else None,
                )
            except Exception:
                logger.exception("Pipeline: не удалось отправить письмо на %s", email)

    finally:
        for path in attachments:
            if os.path.exists(path):
                os.remove(path)


def _status_from_message(message: str) -> str:
    if message == "Purchase updated":
        return "updated"
    if message == "Purchase created":
        return "created"
    return "skipped"


async def _save_purchase_record(purchase: Notice, registration_numbers_dict_per_day: dict, record_semaphore: asyncio.Semaphore) -> None:
    async with record_semaphore:
        try:
            # Ищем заявку

            # Если не находим создаем

            # Если находим обновляем

            # Если при создании произошла ошибка пропускаем и пишем в лог

            data = await api_datum_query(token=settings.system_token, endpoint="put_purchase", **purchase.model_dump(mode = "json"))
            message = data.get("message")
            status = _status_from_message(message)

            if status == "updated":
                logger.info(
                    "Pipeline: обновлена закупка через API | reg=%s | guid=%s",
                    purchase.registration_number,
                    purchase.guid,
                )
            elif status == "created":
                logger.info(
                    "Pipeline: создана новая закупка через API | reg=%s | guid=%s",
                    purchase.registration_number,
                    purchase.guid,
                )
            else:
                logger.warning(
                    "Pipeline: неизвестный ответ API | reg=%s | guid=%s | message=%s",
                    purchase.registration_number,
                    purchase.guid,
                    message,
                )

            registration_numbers_dict_per_day[purchase.filter_type_name][purchase.region_number][status].append(
                purchase.registration_number
            )

        except Exception:
            logger.exception(
                "Pipeline: ошибка сохранения закупки | reg=%s | guid=%s",
                purchase.registration_number,
                purchase.guid,
            )

            filter_type_name = purchase.filter_type_name
            region_number = purchase.region_number
            registration_number = purchase.registration_number
            if filter_type_name and region_number and registration_number:
                registration_numbers_dict_per_day[filter_type_name][region_number][
                    "skipped"
                ].append(registration_number)
            else:
                logger.error(
                    "Pipeline: не удалось учесть запись как skipped — нет ключей | purchase=%s",
                    purchase,
                )


async def _save_protocol_record(protocol: Notice, registration_numbers_dict_per_day: dict, record_semaphore: asyncio.Semaphore) -> None:
    async with record_semaphore:
        try:
            database_answer = await api_datum_query(
                token=settings.system_token,
                endpoint="update_purchase",
                registration_number=protocol.registration_number,
                result_info=protocol.result_info,
                documents_list=protocol.documents_list,
                publication_datetime=protocol.publication_datetime,
            )

            if database_answer.get("message") == "Purchase not found" or database_answer == {}:
                try:
                    data = await api_datum_query(
                        token=settings.system_token, endpoint="put_purchase", **protocol.model_dump(mode = "json")
                    )
                    message = data.get("message")
                    status = _status_from_message(message)

                    if status == "updated":
                        logger.info(
                            "Pipeline: обновлена закупка через API | reg=%s | guid=%s",
                            protocol.registration_number,
                            protocol.guid,
                        )
                    elif status == "created":
                        logger.info(
                            "Pipeline: создана новая закупка через API | reg=%s | guid=%s",
                            protocol.registration_number,
                            protocol.guid,
                        )
                    else:
                        logger.warning(
                            "Pipeline: неизвестный ответ API | reg=%s | guid=%s | message=%s",
                            protocol.registration_number,
                            protocol.guid,
                            message,
                        )

                    registration_numbers_dict_per_day[protocol.filter_type_name][protocol.region_number][
                        status
                    ].append(protocol.registration_number)

                    logger.info("Create purchase from protocol response.json | %s", data)
                    logger.info(
                        "Pipeline: протокол создал закупку | reg=%s",
                        protocol.registration_number,
                    )
                except Exception:
                    logger.exception(
                        "Pipeline: ошибка сохранения закупки | reg=%s | guid=%s",
                        protocol.registration_number,
                        protocol.guid,
                    )
                    filter_type_name = protocol.filter_type_name
                    region_number = protocol.region_number
                    registration_number = protocol.registration_number
                    if filter_type_name and region_number and registration_number:
                        registration_numbers_dict_per_day[filter_type_name][region_number][
                            "skipped"
                        ].append(registration_number)
                    else:
                        logger.error(
                            "Pipeline: не удалось учесть запись как skipped — нет ключей | protocol=%s",
                            protocol,
                        )
                # Закупки не было — создали через put_purchase и уже учли статус выше.
                # Дальше идёт ветка "update_purchase прошёл успешно", сюда она не относится.
                return

            registration_numbers_dict_per_day[protocol.filter_type_name][protocol.region_number][
                "updated"
            ].append(protocol.registration_number)

            logger.info("Update response.json | %s", database_answer)
            logger.info(
                "Pipeline: протокол обновил закупку | reg=%s",
                protocol.registration_number,
            )

        except Exception:
            logger.exception(
                "Pipeline: ошибка update_purchase для протокола | reg=%s",
                protocol.registration_number,
            )
            filter_type_name = protocol.filter_type_name
            region_number = protocol.region_number
            registration_number = protocol.registration_number
            if filter_type_name and region_number and registration_number:
                registration_numbers_dict_per_day[filter_type_name][region_number][
                    "skipped"
                ].append(registration_number)
            else:
                logger.error(
                    "Pipeline: не удалось учесть запись как skipped — нет ключей | purchase=%s",
                    protocol,
                )


async def _process_purchases_for_region(
    region,
    date_str: str,
    filter_number: int,
    registration_numbers_dict_per_day: dict,
    archive_semaphore: asyncio.Semaphore,
    record_semaphore: asyncio.Semaphore,
    attempts: int = 3
) -> None:
    result_purchases = {}

    for attempt in range(1, attempts + 1):
        try:
            result_purchases = await get_docs_by_region(
                org_region=region,
                document_type="purchaseNotice",
                exact_date=date_str,
                subsystem_type="RI223",
            )
            break
        except Exception:
            logger.exception(
                "Pipeline: не удалось получить список закупок | region=%s | попытка %s/%s",
                region, attempt, attempts,
            )
            if attempt == attempts:
                return
            await asyncio.sleep(attempt * 2)

    archive_urls_purchases = result_purchases.get("archive_urls", [])

    async def _handle_archive(archive_url: str) -> None:
        async with archive_semaphore:
            try:
                zip_path_purchases = await download_archive_from_result(archive_url)

                if zip_path_purchases is None:
                    logger.warning("Pipeline: скачивание архива закупок не удалось | region=%s | archive_url=%s", region, archive_url)
                    return

                logger.info("Pipeline: архив закупок скачан: %s", zip_path_purchases)

                purchases = await parse_zip_archive_purchases(zip_path_purchases, region, filter_number)

                logger.info("Pipeline: после фильтров закупок: %s", len(purchases))

                await asyncio.gather(*[
                    _save_purchase_record(purchase, registration_numbers_dict_per_day, record_semaphore)
                    for purchase in purchases
                ])
            except Exception:
                logger.exception(
                    "Pipeline: ошибка обработки архива закупок | region=%s | archive_url=%s",
                    region,
                    archive_url,
                )

    await asyncio.gather(*[_handle_archive(url) for url in archive_urls_purchases])


async def _process_protocols_for_region(
    region,
    date_str: str,
    filter_number: int,
    registration_numbers_dict_per_day: dict,
    archive_semaphore: asyncio.Semaphore,
    record_semaphore: asyncio.Semaphore,
    attempts: int = 3
) -> None:
    result_protocols = {}

    for attempt in range(1, attempts + 1):
        try:
            result_protocols = await get_docs_by_region(
                org_region=region,
                document_type="purchaseProtocol",
                exact_date=date_str,
                subsystem_type="RI223",
            )
            break
        except Exception:
            logger.exception(
                "Pipeline: не удалось получить список закупок | region=%s | попытка %s/%s",
                region, attempt, attempts,
            )
            if attempt == attempts:
                return
            await asyncio.sleep(attempt * 2)

    archive_urls_protocols = result_protocols.get("archive_urls", [])

    async def _handle_archive(archive_url: str) -> None:
        async with archive_semaphore:
            try:
                zip_path_protocols = await download_archive_from_result(archive_url)

                if zip_path_protocols is None:
                    logger.info("Pipeline: скачивание архива не удалось")
                    return

                logger.info("Pipeline: архив протоколов скачан: %s", zip_path_protocols)

                protocols = await parse_zip_archive_protocols(zip_path_protocols, region, filter_number)

                logger.info("Pipeline: после фильтров протоколов: %s", len(protocols))

                await asyncio.gather(*[
                    _save_protocol_record(protocol, registration_numbers_dict_per_day, record_semaphore)
                    for protocol in protocols
                ])
            except Exception:
                logger.exception(
                    "Pipeline: ошибка обработки архива протоколов | region=%s | archive_url=%s",
                    region,
                    archive_url,
                )

    await asyncio.gather(*[_handle_archive(url) for url in archive_urls_protocols])

async def process_period(date_from: datetime, date_to: datetime, filter_number: int):
    logger.info(
        "Processing period of time from %s to %s have started",
        date_from.strftime("%Y-%m-%d"),
        date_to.strftime("%Y-%m-%d"),
    )
    # Строго по одному дню — предсказуемая нагрузка на портал закупок и наш
    # API. process_day и так параллелит регионы/архивы внутри себя.
    while date_from < date_to:
        await process_day(date_from.strftime("%Y-%m-%d"), filter_number=filter_number)
        date_from += timedelta(days=1)

async def process_day(date_str: str, filter_number=0) -> Dict:
    logger.info("Pipeline: обработка даты %s", date_str)

    registration_numbers_dict_per_day = {}

    for filter_name in REGIONS_OF_THE_FILTERS.keys():
        registration_numbers_dict_per_day[filter_name] = {}
        for region_code in REGIONS_OF_THE_FILTERS[filter_name]:
            registration_numbers_dict_per_day[filter_name][region_code] = {
                "created": [],
                "updated": [],
                "skipped": []
            }

    region_semaphore = asyncio.Semaphore(REGION_CONCURRENCY)
    archive_semaphore = asyncio.Semaphore(ARCHIVE_CONCURRENCY)
    record_semaphore = asyncio.Semaphore(RECORD_CONCURRENCY)

    async def _process_region(region) -> None:
        async with region_semaphore:
            try:
                # Закупки и протоколы одного региона не пересекаются по
                # данным — тянем параллельно.
                await asyncio.gather(
                    _process_purchases_for_region(
                        region, date_str, filter_number, registration_numbers_dict_per_day,
                        archive_semaphore, record_semaphore,
                    ),
                    _process_protocols_for_region(
                        region, date_str, filter_number, registration_numbers_dict_per_day,
                        archive_semaphore, record_semaphore,
                    ),
                )
            except Exception:
                logger.exception("Pipeline: не удалось обработать регион %s", region)

    await asyncio.gather(*[_process_region(region) for region in ALL_REGION_CODES])

    return registration_numbers_dict_per_day