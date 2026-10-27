import asyncio
from datetime import datetime, date
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Side, Border
from openpyxl.utils import get_column_letter
import os
import pandas as pd
from app.backend.email.functions import send_email
from app.settings import logger

def _build_analysis_workbook(rows, analysis_path) -> None:
    """Синхронная сборка Excel-файла (была телом create_analysis).
    Вызывается через asyncio.to_thread, чтобы не блокировать event loop."""

    rows.sort(
        key=lambda r: r["Дата публикации"] or datetime.min,
        reverse=True,
    )

    rows = rows[:50000]

    df = pd.DataFrame(rows)
    df.to_excel(analysis_path, index=False)

    wb = load_workbook(analysis_path)
    ws = wb.active

    header_fill = PatternFill("solid", fgColor="366092")
    header_font = Font(color="FFFFFF", bold=True)
    cell_font = Font(size=10)
    align = Alignment(wrap_text=True, vertical="center")
    thin_side = Side(style="thin")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for col in range(1, ws.max_column + 1):
        c = ws.cell(1, col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = align
        c.border = border

    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 30
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = cell_font
            cell.alignment = align
            cell.border = border

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = max(
            len(str(ws.cell(r, col).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[letter].width = min(max_len + 2, 50)

    last = ws.max_row + 2
    ws.cell(last, 1, "Данные по закупкам, лотам и позициям")
    ws.merge_cells(
        start_row=last,
        start_column=1,
        end_row=last,
        end_column=ws.max_column,
    )

    footnote_cell = ws.cell(last, 1)
    footnote_cell.font = Font(italic=True, size=9, color="555555")
    footnote_cell.alignment = Alignment(horizontal="center")
    footnote_cell.border = Border(top=Side(style="thin", color="AAAAAA"))

    wb.save(analysis_path)


async def create_analysis(rows, analysis_path):
    try:
        await asyncio.to_thread(_build_analysis_workbook, rows, analysis_path)
    except Exception:
        logger.exception("Pipeline: ошибка при создании анализа | path=%s", analysis_path)


async def send_analysis(date_of_processing_notices: date, rows: list, emails: list, created: int, updated: int, skipped: int, extra_rows: list = None, path_name_extra="all_purchases.xlsx") -> None:
    date_str = date_of_processing_notices.strftime('%d.%m.%Y')

    logger.info(f"""При обработке закупок за {date_str}\n
    Новых заявок добавлено: {created}\n
    Заявок обновлено: {updated}\n
    Пропущено: {skipped}
    """)

    if created == 0 and updated == 0 and skipped == 0:
        return

    html_content = f"""
                    <html><body style="font-family:Arial;">
                    <h2 style="color:#2E86C1;">Уведомление о заявках с госзакупок</h2>
                    <p>Новых заявок добавлено: <b style="color:#E74C3C;font-size:18px;">{created}</b></p>
                    <p>Заявок обновлено: <b style="color:#F39C12;font-size:18px;">{updated}</b></p>
                    <p>Пропущено: <b style="color:#7F8C8D;font-size:18px;">{skipped}</b></p>
                    <hr><p style="color:#888;font-size:12px;">Это письмо сформировано автоматически, отвечать на него не нужно. 
                    В excel-таблицах отображается максимум 50 000 строк-заявок.</p>
                    </body></html>
                    """

    attachments = []

    analysis_path = "analysis.xlsx"

    await create_analysis(rows, analysis_path=analysis_path)

    attachments.append(analysis_path)

    if extra_rows:

        await create_analysis(extra_rows, analysis_path=path_name_extra)
        attachments.append(path_name_extra)

    subject = f"Заявки с госзакупок за {date_str}"

    try:

        for email in emails:
            try:
                await send_email(
                    email,
                    subject,
                    html_content,
                    attachments=attachments if (created or updated) else None,
                )
            except Exception:
                logger.exception("Pipeline: не удалось отправить письмо на %s", email)

    finally:
        for path in attachments:
            if os.path.exists(path):
                os.remove(path)